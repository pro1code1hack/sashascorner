"""`cafeops import-finance`: the finance workbook's Daily Sales, Expenses and Director
Account sheets into the finance tables. Finance spec 3.6; DECISIONS.md 4.

Rules, each with a reason:

* **Every row is LEGACY_WORKBOOK** with a `source_ref` naming its sheet and row, so a
  wrong figure is traceable and the import is idempotent: a second run matches by
  `source_ref` (or the (date, method, source) key for takings) and changes nothing.
* **A row somebody edited in the app is never overwritten.** Expenses with an
  `updated_by`, trading days whose `source` became MANUAL, and soft-deleted expenses are
  reported and left alone.
* **Card figures up to 31 Mar 2026 are bank deposits** (README row 36): basis
  BANK_DEPOSIT. April is till data: TILL.
* **The workbook's "Square cash" column holds Just Eat money** on 7 rows (£114.92,
  README row 37). Just Eat never paid cash (DECISIONS 4), so those amounts become Just
  Eat monthly statements (gross only; commission and ads unreported, never zero) and
  no cash row is written for them.
* **Renames** (DECISIONS 4): "Square fees" -> "Card fees". Square cash and Own cash are
  one figure, "Cash" (DECISIONS 26): both columns sum into a single CASH row.
* **Drawings are stored once.** Each DRAWINGS expense creates its director's-account
  mirror; the Director Account sheet's Drawings rows are matched to those mirrors by
  (date, description, amount) and linked, not inserted a second time.
* Nothing from the design's demo data (invented Deliveroo months, fake cash counts, a
  fake short payout) is imported. There is none in the workbook to import.
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from cafeops.db.models.finance import (
    ChannelStatement,
    DirectorEntry,
    Expense,
    ExpenseCategory,
    TradingDay,
)
from cafeops.db.models.payment import PaymentDay
from cafeops.domain.enums import (
    ChannelSourceKind,
    DirectorEntryType,
    ExpenseKind,
    ExpenseMethod,
    FinanceSource,
    PaymentBasis,
    PaymentMethod,
    PaymentSourceKind,
    SalesChannelName,
)
from cafeops.services.finance.expenses import (
    NEEDS_REVIEW_RE,
    kind_from_notes,
    sync_director_mirror,
)

__all__ = ["FinanceImportReport", "import_finance"]

#: README row 36: "Card sales for Sep-Mar are from Mettle Square deposits".
LAST_DEPOSIT_BASIS_DAY = date(2026, 3, 31)
JUST_EAT_NOTE = re.compile(r"^\s*Just Eat:\s*£\s*([0-9]+(?:\.[0-9]{1,2})?)", re.IGNORECASE)
CATEGORY_RENAMES = {"Square fees": "Card fees"}
METHODS = {
    "card": ExpenseMethod.CARD,
    "bank transfer": ExpenseMethod.BANK_TRANSFER,
    "direct debit": ExpenseMethod.DIRECT_DEBIT,
    "standing order": ExpenseMethod.STANDING_ORDER,
    "cash": ExpenseMethod.CASH,
    "cash withdrawal": ExpenseMethod.CASH_WITHDRAWAL,
}
DIRECTOR_TYPES = {
    "capital injection": DirectorEntryType.CAPITAL_INJECTION,
    "loan to company": DirectorEntryType.LOAN_TO_COMPANY,
    "drawings": DirectorEntryType.DRAWINGS,
    "repayment": DirectorEntryType.REPAYMENT,
}


@dataclass
class SheetCounts:
    read: int = 0
    inserted: int = 0
    updated: int = 0
    unchanged: int = 0
    skipped_edited: int = 0

    def line(self) -> str:
        return (
            f"{self.read} read, +{self.inserted} new, ~{self.updated} updated, "
            f"{self.unchanged} unchanged, {self.skipped_edited} left alone (edited in the app)"
        )


@dataclass
class FinanceImportReport:
    workbook: Path
    days: SheetCounts = field(default_factory=SheetCounts)
    takings_rows: SheetCounts = field(default_factory=SheetCounts)
    expenses: SheetCounts = field(default_factory=SheetCounts)
    director: SheetCounts = field(default_factory=SheetCounts)
    just_eat: SheetCounts = field(default_factory=SheetCounts)
    card_pence: int = 0
    #: Till cash + own cash from the workbook, stored as one CASH figure (DECISIONS 26).
    cash_pence: int = 0
    just_eat_pence: int = 0
    just_eat_rows: int = 0
    expense_total_pence: int = 0
    expense_by_kind: dict[str, tuple[int, int]] = field(default_factory=dict)
    needs_review: int = 0
    director_in_pence: int = 0
    director_out_pence: int = 0
    drawings_matched: int = 0
    drawings_rows: int = 0
    drawings_unmatched: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    refused: list[str] = field(default_factory=list)

    @property
    def changed(self) -> int:
        return sum(
            c.inserted + c.updated
            for c in (self.days, self.takings_rows, self.expenses, self.director, self.just_eat)
        )


# --------------------------------------------------------------------------
# cell helpers
# --------------------------------------------------------------------------


def _pence(value: Any, where: str) -> int | None:
    """A £ cell to integer pence, exactly. Refuses a third decimal place."""
    if value is None or value == "":
        return None
    if isinstance(value, str) and value.startswith("="):
        raise ValueError(f"{where}: formula without a cached value ({value})")
    try:
        d = Decimal(str(value).replace("£", "").replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError(f"{where}: {value!r} is not money") from exc
    pence = d * 100
    if pence != pence.to_integral_value():
        raise ValueError(f"{where}: {value!r} has more than two decimal places")
    return int(pence)


def _day(value: Any, where: str) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip()[:10])
    except ValueError as exc:
        raise ValueError(f"{where}: {value!r} is not a date") from exc


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _header(ws: Any, row: int, expect: dict[int, str]) -> None:
    """Refuse a sheet whose header row moved: guessing columns moves money."""
    cells = [c.value for c in ws[row]]
    for col, needle in expect.items():
        got = str(cells[col - 1] or "")
        if needle.lower() not in got.lower():
            raise ValueError(
                f"{ws.title}!{row}: column {col} header is {got!r}, expected {needle!r}. "
                "The sheet layout changed; nothing was imported."
            )


# --------------------------------------------------------------------------
# Daily Sales
# --------------------------------------------------------------------------


def _upsert_payment(
    session: Session,
    rep: FinanceImportReport,
    day: date,
    method: PaymentMethod,
    pence: int,
    basis: PaymentBasis,
    ref: str,
    transactions: int | None,
) -> None:
    rep.takings_rows.read += 1
    row = session.scalars(
        select(PaymentDay).where(
            PaymentDay.business_date == day,
            PaymentDay.method == method,
            PaymentDay.source == PaymentSourceKind.LEGACY_WORKBOOK,
        )
    ).one_or_none()
    if row is None:
        session.add(
            PaymentDay(
                business_date=day,
                method=method,
                gross_pence=pence,
                basis=basis,
                transactions=transactions,
                source=PaymentSourceKind.LEGACY_WORKBOOK,
                source_ref=ref,
                notes="finance workbook, Daily Sales",
            )
        )
        rep.takings_rows.inserted += 1
        return
    if (row.gross_pence, row.basis, row.transactions, row.source_ref) == (
        pence,
        basis,
        transactions,
        ref,
    ):
        rep.takings_rows.unchanged += 1
        return
    row.gross_pence, row.basis, row.transactions, row.source_ref = pence, basis, transactions, ref
    rep.takings_rows.updated += 1


def _import_daily_sales(session: Session, wb: Any, rep: FinanceImportReport) -> None:
    ws = wb["Daily Sales"]
    _header(
        ws,
        5,
        {2: "Date", 4: "card", 5: "cash", 6: "Own cash", 7: "override", 9: "Trans", 11: "Notes"},
    )
    just_eat: dict[date, list[tuple[int, str]]] = defaultdict(list)
    for n, cells in enumerate(ws.iter_rows(min_row=6, values_only=True), start=6):
        b, d_card, e_cash, f_own, g_override, i_tx, k_note = (
            cells[1],
            cells[3],
            cells[4],
            cells[5],
            cells[6],
            cells[8],
            cells[10],
        )
        where = f"Daily Sales!R{n}"
        try:
            day = _day(b, where)
            card = _pence(d_card, f"Daily Sales!D{n}")
            cash = _pence(e_cash, f"Daily Sales!E{n}")
            override = _pence(g_override, f"Daily Sales!G{n}")
            own_auto = _pence(f_own, f"Daily Sales!F{n}") if not isinstance(f_own, str) else None
        except ValueError as exc:
            rep.refused.append(str(exc))
            continue
        tx = int(i_tx) if isinstance(i_tx, (int, float)) and i_tx > 0 else None
        note = _text(k_note)
        own = override if override else own_auto
        if day is None or not any([card, cash, override, tx, note]):
            continue
        rep.days.read += 1

        # The trading-day row: note and provenance.
        td = session.scalars(
            select(TradingDay).where(TradingDay.business_date == day)
        ).one_or_none()
        if td is None:
            session.add(
                TradingDay(
                    business_date=day,
                    note=note,
                    source=FinanceSource.LEGACY_WORKBOOK,
                    source_ref=where,
                )
            )
            rep.days.inserted += 1
        elif td.source is not FinanceSource.LEGACY_WORKBOOK:
            rep.days.skipped_edited += 1
        elif (td.note, td.source_ref) == (note, where):
            rep.days.unchanged += 1
        else:
            td.note, td.source_ref = note, where
            rep.days.updated += 1

        if card:
            basis = (
                PaymentBasis.BANK_DEPOSIT if day <= LAST_DEPOSIT_BASIS_DAY else PaymentBasis.TILL
            )
            _upsert_payment(
                session, rep, day, PaymentMethod.CARD, card, basis, f"Daily Sales!D{n}", tx
            )
            rep.card_pence += card
            tx = None  # the count rides on one row only
        # One cash figure per day (DECISIONS 26): the workbook's till cash (E) and its own
        # cash (G override, else F auto) are written as ONE CASH row. Own cash is never
        # stored apart any more; nothing is lost, it is summed into the day's Cash.
        till_cash = 0
        if cash:
            m = JUST_EAT_NOTE.match(note or "")
            if m:
                noted = _pence(m.group(1), f"Daily Sales!K{n}")
                if noted != cash:
                    rep.warnings.append(
                        f"Daily Sales!R{n}: note says Just Eat £{m.group(1)} but the cash "
                        f"column holds {cash}p; the column figure is used"
                    )
                just_eat[day.replace(day=1)].append((cash, f"E{n}"))
                rep.just_eat_pence += cash
                rep.just_eat_rows += 1
            else:
                till_cash = cash
        own_cash = own or 0
        if till_cash or own_cash:
            refs = ([f"E{n}"] if till_cash else []) + (
                [f"{'G' if override else 'F'}{n}"] if own_cash else []
            )
            _upsert_payment(
                session,
                rep,
                day,
                PaymentMethod.CASH,
                till_cash + own_cash,
                PaymentBasis.TILL,
                "Daily Sales!" + "+".join(refs),
                tx,
            )
            rep.cash_pence += till_cash + own_cash
            tx = None

    for month, parts in sorted(just_eat.items()):
        rep.just_eat.read += 1
        gross = sum(p for p, _ in parts)
        ref = "Daily Sales!" + ",".join(r for _, r in parts)
        st = session.scalars(
            select(ChannelStatement).where(
                ChannelStatement.channel == SalesChannelName.JUST_EAT,
                ChannelStatement.month == month,
            )
        ).one_or_none()
        if st is None:
            session.add(
                ChannelStatement(
                    channel=SalesChannelName.JUST_EAT,
                    month=month,
                    gross_pence=gross,
                    commission_pence=None,
                    ad_spend_pence=None,
                    source=ChannelSourceKind.LEGACY_WORKBOOK,
                    source_ref=ref,
                    notes=(
                        "Just Eat money the workbook kept in its cash column. Commission "
                        "and ads were not recorded."
                    ),
                )
            )
            rep.just_eat.inserted += 1
        elif st.source is not ChannelSourceKind.LEGACY_WORKBOOK:
            rep.just_eat.skipped_edited += 1
        elif (st.gross_pence, st.source_ref) == (gross, ref):
            rep.just_eat.unchanged += 1
        else:
            st.gross_pence, st.source_ref = gross, ref
            rep.just_eat.updated += 1


# --------------------------------------------------------------------------
# Expenses
# --------------------------------------------------------------------------


def _import_expenses(session: Session, wb: Any, rep: FinanceImportReport) -> None:
    ws = wb["Expenses"]
    _header(
        ws,
        5,
        {
            2: "Date",
            3: "Category",
            4: "Vendor",
            5: "Amount",
            6: "Payment",
            7: "Receipt",
            8: "Notes",
        },
    )
    cats = {c.name: c for c in session.scalars(select(ExpenseCategory))}
    existing = {
        e.source_ref: e
        for e in session.scalars(
            select(Expense).where(Expense.source == FinanceSource.LEGACY_WORKBOOK)
        )
    }
    by_kind: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for n, cells in enumerate(ws.iter_rows(min_row=6, max_row=500, values_only=True), start=6):
        b, c, d, e, f, g, h = cells[1:8]
        if b is None and d is None and e is None:
            continue
        ref = f"Expenses!R{n}"
        try:
            paid_on = _day(b, f"Expenses!B{n}")
            amount = _pence(e, f"Expenses!E{n}")
        except ValueError as exc:
            rep.refused.append(str(exc))
            continue
        desc = _text(d)
        cat_name = CATEGORY_RENAMES.get(_text(c) or "", _text(c) or "")
        if paid_on is None or amount is None or amount <= 0 or not desc:
            rep.refused.append(f"{ref}: needs a date, a description and an amount over zero")
            continue
        cat = cats.get(cat_name)
        if cat is None:
            rep.refused.append(f"{ref}: category {cat_name!r} is not one of the 13")
            continue
        rep.expenses.read += 1
        notes = _text(h)
        raw_method = _text(f)
        method = METHODS.get((raw_method or "").lower()) if raw_method else None
        if raw_method and method is None:
            method = ExpenseMethod.OTHER
            notes = f"{notes} (paid by: {raw_method})" if notes else f"(paid by: {raw_method})"
        kind = kind_from_notes(notes)
        needs_review = bool(NEEDS_REVIEW_RE.search(notes or ""))
        receipt = bool(g) and str(g).strip().lower() not in {"no", "n", "0", "false"}
        rep.expense_total_pence += amount
        by_kind[kind.value][0] += 1
        by_kind[kind.value][1] += amount
        rep.needs_review += int(needs_review)

        row = existing.get(ref)
        if row is None:
            row = Expense(
                paid_on=paid_on,
                category=cat,
                description=desc[:300],
                amount_pence=amount,
                method=method,
                kind=kind,
                has_receipt=receipt,
                notes=notes,
                needs_review=needs_review,
                source=FinanceSource.LEGACY_WORKBOOK,
                source_ref=ref,
            )
            session.add(row)
            session.flush()
            sync_director_mirror(session, row)
            rep.expenses.inserted += 1
            continue
        if row.updated_by is not None or row.deleted_at is not None:
            rep.expenses.skipped_edited += 1
            continue
        wanted = (paid_on, cat.id, desc[:300], amount, method, kind, receipt, notes, needs_review)
        current = (
            row.paid_on,
            row.category_id,
            row.description,
            row.amount_pence,
            row.method,
            row.kind,
            row.has_receipt,
            row.notes,
            row.needs_review,
        )
        if wanted == current:
            rep.expenses.unchanged += 1
            continue
        (
            row.paid_on,
            row.category_id,
            row.description,
            row.amount_pence,
            row.method,
            row.kind,
            row.has_receipt,
            row.notes,
            row.needs_review,
        ) = wanted
        session.flush()
        sync_director_mirror(session, row)
        rep.expenses.updated += 1
    rep.expense_by_kind = {k: (v[0], v[1]) for k, v in by_kind.items()}


# --------------------------------------------------------------------------
# Director Account
# --------------------------------------------------------------------------


def _import_director(session: Session, wb: Any, rep: FinanceImportReport) -> None:
    ws = wb["Director Account"]
    _header(ws, 5, {2: "Date", 3: "Type", 4: "Description", 5: "IN", 6: "OUT", 7: "Notes"})
    by_ref = {
        d.source_ref: d
        for d in session.scalars(select(DirectorEntry).where(DirectorEntry.source_ref.is_not(None)))
    }
    # Unclaimed drawings mirrors, for matching the sheet's Drawings rows.
    mirrors = list(
        session.scalars(
            select(DirectorEntry)
            .join(Expense, Expense.id == DirectorEntry.expense_id)
            .where(
                DirectorEntry.deleted_at.is_(None), Expense.source == FinanceSource.LEGACY_WORKBOOK
            )
            .order_by(DirectorEntry.id)
        )
    )
    claimed = {d.id for d in mirrors if d.source_ref and d.source_ref.startswith("Director")}

    for n, cells in enumerate(ws.iter_rows(min_row=6, values_only=True), start=6):
        b, c, d, e, f, g = cells[1:7]
        if str(d or "").strip().upper().startswith("TOTAL"):
            break
        if b is None and c is None:
            continue
        ref = f"Director Account!R{n}"
        try:
            day = _day(b, f"Director Account!B{n}")
            in_p = _pence(e, f"Director Account!E{n}") or 0
            out_p = _pence(f, f"Director Account!F{n}") or 0
        except ValueError as exc:
            rep.refused.append(str(exc))
            continue
        type_ = DIRECTOR_TYPES.get((_text(c) or "").lower())
        desc = _text(d)
        if day is None or type_ is None or not desc or (in_p > 0) == (out_p > 0):
            rep.refused.append(f"{ref}: needs a date, a known type, a description and one amount")
            continue
        rep.director.read += 1
        rep.director_in_pence += in_p
        rep.director_out_pence += out_p
        notes = _text(g)

        if type_ is DirectorEntryType.DRAWINGS:
            rep.drawings_rows += 1
            linked = by_ref.get(ref)
            if linked is not None and linked.expense_id is not None:
                rep.drawings_matched += 1
                rep.director.unchanged += 1
                continue
            match = next(
                (
                    m
                    for m in mirrors
                    if m.id not in claimed
                    and m.entry_date == day
                    and m.description == desc
                    and m.out_pence == out_p
                ),
                None,
            )
            if match is not None:
                claimed.add(match.id)
                match.source_ref = ref
                match.notes = notes
                rep.drawings_matched += 1
                rep.director.updated += 1
                continue
            rep.drawings_unmatched.append(f"{ref}: {day} {desc} {out_p}p")

        row = by_ref.get(ref)
        if row is None:
            session.add(
                DirectorEntry(
                    entry_date=day,
                    type=type_,
                    description=desc[:300],
                    in_pence=in_p,
                    out_pence=out_p,
                    notes=notes,
                    source=FinanceSource.LEGACY_WORKBOOK,
                    source_ref=ref,
                )
            )
            rep.director.inserted += 1
        elif row.updated_by is not None or row.deleted_at is not None:
            rep.director.skipped_edited += 1
        elif (
            row.entry_date,
            row.type,
            row.description,
            row.in_pence,
            row.out_pence,
            row.notes,
        ) == (day, type_, desc[:300], in_p, out_p, notes):
            rep.director.unchanged += 1
        else:
            (row.entry_date, row.type, row.description, row.in_pence, row.out_pence, row.notes) = (
                day,
                type_,
                desc[:300],
                in_p,
                out_p,
                notes,
            )
            rep.director.updated += 1


def import_finance(session: Session, workbook: Path) -> FinanceImportReport:
    """Read the three sheets and write them. Does not commit: the caller decides."""
    wb = load_workbook(workbook, data_only=True, read_only=False)
    rep = FinanceImportReport(workbook=workbook)
    for sheet in ("Daily Sales", "Expenses", "Director Account"):
        if sheet not in wb.sheetnames:
            raise ValueError(f"{workbook.name}: no {sheet!r} sheet; is this the finance workbook?")
    _import_daily_sales(session, wb, rep)
    session.flush()
    _import_expenses(session, wb, rep)
    session.flush()
    _import_director(session, wb, rep)
    session.flush()
    if ExpenseKind.DRAWINGS.value in rep.expense_by_kind and rep.drawings_rows:
        n_exp = rep.expense_by_kind[ExpenseKind.DRAWINGS.value][0]
        if n_exp != rep.drawings_rows:
            rep.warnings.append(
                f"{n_exp} drawings expenses but {rep.drawings_rows} Drawings rows on the "
                "Director Account sheet"
            )
    return rep


def _gbp(pence: int) -> str:
    sign = "-" if pence < 0 else ""
    whole, part = divmod(abs(pence), 100)
    return f"{sign}£{whole:,}.{part:02d}"


def report_lines(rep: FinanceImportReport) -> list[str]:
    out = [
        f"workbook: {rep.workbook}",
        f"Daily Sales days      {rep.days.line()}",
        f"  takings rows        {rep.takings_rows.line()}",
        f"  card {_gbp(rep.card_pence)} · cash {_gbp(rep.cash_pence)}",
        f"  Just Eat months     {rep.just_eat.line()}",
        f"  Just Eat money      {_gbp(rep.just_eat_pence)} over {rep.just_eat_rows} rows "
        "(imported as Just Eat takings, not cash)",
        f"Expenses              {rep.expenses.line()}",
        f"  total {_gbp(rep.expense_total_pence)} · needs a look {rep.needs_review}",
    ]
    for kind, (n, pence) in sorted(rep.expense_by_kind.items()):
        out.append(f"  {kind:<10} {n:>4} rows {_gbp(pence):>12}")
    out += [
        f"Director Account      {rep.director.line()}",
        f"  in {_gbp(rep.director_in_pence)} · out {_gbp(rep.director_out_pence)} · "
        f"drawings matched to expenses {rep.drawings_matched}/{rep.drawings_rows}",
    ]
    out += [f"  unmatched drawing: {u}" for u in rep.drawings_unmatched]
    out += [f"WARNING {w}" for w in rep.warnings]
    out += [f"REFUSED {r}" for r in rep.refused]
    out.append(f"rows changed: {rep.changed}")
    return out


def check_against_workbook(session: Session, workbook: Path) -> list[str]:
    """Recompute months and compare with the workbook's cached Monthly P&L.

    Differences are expected and explained: the workbook counts Just Eat as cash, and
    its "Other" line holds capital and drawings, which the P&L here keeps out.
    """
    from cafeops.services.finance.periods import period_figures

    wb = load_workbook(workbook, data_only=True)
    ws = wb["Monthly P&L"]
    year = next(
        (r[2].value for r in ws.iter_rows(min_row=1, max_row=8) if r[1].value == "Year:"), None
    )
    if not isinstance(year, int):
        return ["check against the workbook skipped: no Year cell on Monthly P&L"]
    labels = {str(r[1].value).strip(): r for r in ws.iter_rows(min_row=5, max_row=45) if r[1].value}
    out = [f"check against the workbook's Monthly P&L ({year}):"]
    for i, name in enumerate(("Jan", "Feb", "Mar", "Apr")):
        month = date(int(year), i + 1, 1)
        fig = period_figures(session, month)

        def cell(label: str, col: int = i) -> int:
            row = labels.get(label)
            v = row[2 + col].value if row is not None else None
            return _pence(v, label) or 0

        je = fig.delivery_gross_pence or 0
        pairs = (
            ("card", cell("Card sales"), fig.card_pence),
            ("cash (+Just Eat)", cell("Square cash sales"), fig.cash_pence + je),
            ("COGS", cell("TOTAL COGS"), fig.stock_bought_pence),
            (
                "Rent",
                cell("Rent"),
                next((c.pence for c in fig.opex_by_category if c.category == "Rent"), 0),
            ),
            (
                "Utilities",
                cell("Utilities"),
                next((c.pence for c in fig.opex_by_category if c.category == "Utilities"), 0),
            ),
            (
                "Other (+capital+drawings)",
                cell("Other"),
                next((c.pence for c in fig.opex_by_category if c.category == "Other"), 0)
                + fig.capital_pence
                + fig.drawings_pence,
            ),
        )
        bits = []
        for label, theirs, ours in pairs:
            mark = "ok" if theirs == ours else f"DIFF {_gbp(ours - theirs)}"
            bits.append(f"{label} {_gbp(theirs)} {mark}")
        out.append(f"  {name} {year}: " + " · ".join(bits))
    return out
